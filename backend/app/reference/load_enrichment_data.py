"""Loads the internal reference workbook into plain text blocks (one per
sheet) for use as enrichment context in Gemini prompts.

Reads sheets and header rows generically -- no hardcoded column names --
so the real reference workbook can be dropped in once provided without any
code changes here.
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

# TODO: replace with real enrichment_data.xlsx once provided
ENRICHMENT_DATA_PATH = Path(__file__).resolve().parent / "enrichment_data.xlsx"


def _format_row(headers: list[str], row: tuple) -> str:
    parts = []
    for header, value in zip(headers, row):
        if value is None or str(value).strip() == "":
            continue
        parts.append(f"{header}: {value}")
    return " | ".join(parts)


def _sheet_to_text(sheet) -> str:
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return ""

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    lines = []
    for row in rows:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        line = _format_row(headers, row)
        if line:
            lines.append(line)
    return "\n".join(lines)


def load_enrichment_data() -> str:
    """Read every sheet of the reference workbook into one text blob, one
    section per sheet. Returns an empty string (with a warning logged) if
    the workbook doesn't exist yet."""
    if not ENRICHMENT_DATA_PATH.exists():
        logger.warning(
            "enrichment_data.xlsx not found at %s; enrichment will run with no reference data",
            ENRICHMENT_DATA_PATH,
        )
        return ""

    workbook = openpyxl.load_workbook(ENRICHMENT_DATA_PATH, read_only=True, data_only=True)
    try:
        blocks = [f"### {name}\n{_sheet_to_text(workbook[name])}" for name in workbook.sheetnames]
        return "\n\n".join(blocks)
    finally:
        workbook.close()
