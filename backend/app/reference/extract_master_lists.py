"""One-time script: workbook master-list sheets -> plain text files.

Run with: python -m app.reference.extract_master_lists
Reads backend/sample_data/workbook.xlsx and writes one .txt file per master
list sheet under app/reference/, with one line per data row built dynamically
from whatever header columns that sheet actually has.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

BACKEND_DIR = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = BACKEND_DIR / "sample_data" / "workbook.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parent

# Sheet name -> output filename. Matched case-insensitively against a
# substring of the workbook's actual sheet names, since sheet names can carry
# incidental whitespace or minor variation.
TARGET_SHEETS = {
    "legal entity master list": "legal_entity_master_list.txt",
    "vendor master list": "vendor_master_list.txt",
    "investor master list": "investor_master_list.txt",
    "project code report": "project_code_report.txt",
}


def _find_sheet_name(workbook_sheet_names: list[str], needle: str) -> str | None:
    needle_normalized = needle.strip().lower()
    for name in workbook_sheet_names:
        if needle_normalized in name.strip().lower():
            return name
    return None


def _format_row(headers: list[str], row: tuple) -> str:
    parts = []
    for header, value in zip(headers, row):
        if value is None or str(value).strip() == "":
            continue
        parts.append(f"{header}: {value}")
    return " | ".join(parts)


def extract_sheet_to_text(sheet, out_path: Path) -> int:
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        out_path.write_text("", encoding="utf-8")
        return 0

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    lines = []
    for row in rows:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        line = _format_row(headers, row)
        if line:
            lines.append(line)

    out_path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    return len(lines)


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Workbook not found at {WORKBOOK_PATH}")

    workbook = openpyxl.load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        for needle, out_filename in TARGET_SHEETS.items():
            sheet_name = _find_sheet_name(workbook.sheetnames, needle)
            out_path = OUTPUT_DIR / out_filename
            if sheet_name is None:
                print(f"WARNING: no sheet matching '{needle}' found; writing empty {out_filename}")
                out_path.write_text("", encoding="utf-8")
                continue
            count = extract_sheet_to_text(workbook[sheet_name], out_path)
            print(f"Wrote {count} rows from sheet '{sheet_name}' -> {out_path.name}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
