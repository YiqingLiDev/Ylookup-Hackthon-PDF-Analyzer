"""Writes the final .xlsx: 18 columns, one row per transaction, with an
embedded screenshot of the source page in each row. Sorted by confidence
ascending so the rows most needing review are at the top.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

COLUMNS = [
    "account_name",
    "account_number",
    "currency",
    "bank_reference",
    "customer_reference",
    "trn_type",
    "value_date",
    "credit_amount",
    "debit_amount",
    "balance",
    "post_date",
    "narrative",
    "company",
    "metric",
    "confidence",
    "source_pdf",
    "page",
    "screenshot",
]

IMAGE_WIDTH_PX = 200
IMAGE_COLUMN_WIDTH = IMAGE_WIDTH_PX / 7.0  # approx openpyxl char-width units
ROW_HEIGHT_PT = 150


def build_workbook(rows: list[dict], output_path: Path) -> Path:
    ordered_rows = sorted(rows, key=lambda r: r["confidence"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Sourceline"

    ws.append(COLUMNS)

    screenshot_col_idx = len(COLUMNS)
    screenshot_col_letter = get_column_letter(screenshot_col_idx)
    ws.column_dimensions[screenshot_col_letter].width = IMAGE_COLUMN_WIDTH

    for i, row in enumerate(ordered_rows, start=2):
        values = [
            row.get("account_name"),
            row.get("account_number"),
            row.get("currency"),
            row.get("bank_reference"),
            row.get("customer_reference"),
            row.get("trn_type"),
            row.get("value_date"),
            row.get("credit_amount"),
            row.get("debit_amount"),
            row.get("balance"),
            row.get("post_date"),
            row.get("narrative"),
            row.get("company"),
            row.get("metric"),
            row.get("confidence"),
            row.get("source_pdf"),
            row.get("page"),
            None,  # screenshot is embedded as an image, not a cell value
        ]
        ws.append(values)
        ws.row_dimensions[i].height = ROW_HEIGHT_PT

        screenshot_path = row.get("screenshot_path")
        if screenshot_path and Path(screenshot_path).exists():
            img = XLImage(screenshot_path)
            aspect_ratio = img.height / img.width
            img.width = IMAGE_WIDTH_PX
            img.height = int(IMAGE_WIDTH_PX * aspect_ratio)
            ws.add_image(img, f"{screenshot_col_letter}{i}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
